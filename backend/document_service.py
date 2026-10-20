"""Ephemeral extraction of supported teaching-document formats."""

from io import BytesIO
import shutil
import subprocess
import tempfile
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation

MAX_EXTRACTED_CHARACTERS = 12000
MAX_UPLOAD_BYTES = 4 * 1024 * 1024
MAX_PDF_PAGES = 40
MAX_PRESENTATION_SLIDES = 80
SUPPORTED_FORMATS = {".pdf", ".pptx", ".doc", ".docx"}


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Extract bounded plain text without retaining the uploaded file."""
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_FORMATS:
        raise ValueError("Upload a PDF, PowerPoint, or Word document.")
    if not file_bytes:
        raise ValueError("The uploaded file is empty.")
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise ValueError("Documents are limited to 4 MB.")

    try:
        if extension == ".pdf":
            text = _extract_pdf(file_bytes)
        elif extension == ".pptx":
            text = _extract_presentation(file_bytes)
        elif extension == ".docx":
            text = _extract_docx(file_bytes)
        else:
            text = _extract_legacy_word(file_bytes)
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("The document could not be read. Check the file and try again.") from error

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not text:
        raise ValueError("No selectable text was found in this file.")
    return text[:MAX_EXTRACTED_CHARACTERS]


def _extract_pdf(file_bytes: bytes) -> str:
    """Extract selectable text from a bounded number of PDF pages."""
    with pymupdf.open(stream=file_bytes, filetype="pdf") as document:
        if len(document) > MAX_PDF_PAGES:
            raise ValueError(f"PDFs are limited to {MAX_PDF_PAGES} pages.")
        return "\n".join(page.get_text() for page in document)


def _extract_presentation(file_bytes: bytes) -> str:
    """Extract text from a bounded number of presentation slides."""
    presentation = Presentation(BytesIO(file_bytes))
    if len(presentation.slides) > MAX_PRESENTATION_SLIDES:
        raise ValueError(f"Presentations are limited to {MAX_PRESENTATION_SLIDES} slides.")
    return "\n".join(
        shape.text for slide in presentation.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    )


def _extract_docx(file_bytes: bytes) -> str:
    """Extract text from paragraphs and tables in a modern Word document."""
    document = Document(BytesIO(file_bytes))
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def _extract_legacy_word(file_bytes: bytes) -> str:
    """Use antiword for legacy .doc files when that executable is installed."""
    executable = shutil.which("antiword")
    if executable is None:
        raise ValueError("Legacy .doc extraction is unavailable in this runtime. Save the document as .docx and retry.")
    with tempfile.NamedTemporaryFile(suffix=".doc") as source_file:
        source_file.write(file_bytes)
        source_file.flush()
        result = subprocess.run(
            [executable, source_file.name],
            capture_output=True,
            check=False,
            timeout=10,
        )
    if result.returncode:
        raise ValueError("The legacy Word document could not be read. Try saving it as .docx.")
    return result.stdout.decode("utf-8", errors="replace")