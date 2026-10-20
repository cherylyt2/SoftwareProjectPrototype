"""Upload parsing success, bounds, and failure behavior."""

from io import BytesIO

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from document_service import extract_text


def test_extracts_docx_paragraph_text():
    """A valid Word document yields its paragraph text."""
    document = Document()
    document.add_paragraph("Cell structure and function")
    source = BytesIO()
    document.save(source)

    assert extract_text("notes.docx", source.getvalue()) == "Cell structure and function"


def test_extracts_pptx_slide_text():
    """A valid presentation yields visible slide text."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    textbox.text = "A useful learning objective"
    source = BytesIO()
    presentation.save(source)

    assert extract_text("lesson.pptx", source.getvalue()) == "A useful learning objective"


def test_rejects_unsupported_empty_and_oversized_files():
    """Invalid upload types and sizes return clear parser errors."""
    try:
        extract_text("lesson.exe", b"data")
        assert False, "Unsupported extensions must be rejected."
    except ValueError as error:
        assert "Upload a PDF" in str(error)
    try:
        extract_text("lesson.pdf", b"")
        assert False, "Empty documents must be rejected."
    except ValueError as error:
        assert "empty" in str(error)
    try:
        extract_text("lesson.docx", b"x" * (4 * 1024 * 1024 + 1))
        assert False, "Oversized documents must be rejected."
    except ValueError as error:
        assert "4 MB" in str(error)


def test_legacy_doc_fails_clearly_without_antiword(monkeypatch):
    """Legacy Word files explain the missing runtime parser rather than crashing."""
    monkeypatch.setattr("document_service.shutil.which", lambda _: None)

    try:
        extract_text("legacy.doc", b"binary content")
        assert False, "Legacy parsing must report its missing executable."
    except ValueError as error:
        assert ".docx" in str(error)